import os
import io
import json
import pandas as pd
# pyrefly: ignore [missing-import]
from fastapi import FastAPI, UploadFile, File, Form, HTTPException
# pyrefly: ignore [missing-import]
from fastapi.middleware.cors import CORSMiddleware
# pyrefly: ignore [missing-import]
from fastapi.responses import StreamingResponse, FileResponse
# pyrefly: ignore [missing-import]
from fastapi.staticfiles import StaticFiles
# pyrefly: ignore [missing-import]
from pydantic import BaseModel
from typing import List, Optional

# Import local backend modules
from backend.config import SESSION, LLM_PROVIDERS
from backend.cleaning.profiler import profile_dataset
from backend.cleaning.statistical import detect_anomalies
from backend.cleaning.recommender import generate_recommendations
from backend.cleaning.engine import apply_cleaning_recommendations
from backend.cleaning.validator import validate_cleaning
from backend.report_generator import generate_pdf_report
from backend.utils.dataio import (
    read_tabular, FileLoadError, sanitize_for_json, df_to_records, safe_number,
)

app = FastAPI(title="Automated Data Cleaning & Preprocessing API")

# Add CORS Middleware to support local developments
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# Request Models
class AnalyzeRequest(BaseModel):
    provider: str
    api_key: Optional[str] = None

class CleanRequest(BaseModel):
    active_cols: Optional[List[str]] = None

@app.get("/api/providers")
def get_providers():
    """Returns list of supported LLM Providers."""
    return {
        "providers": LLM_PROVIDERS,
        "env_openai": bool(os.getenv("OPENAI_API_KEY")),
        "env_anthropic": bool(os.getenv("ANTHROPIC_API_KEY"))
    }

@app.post("/api/upload")
async def upload_dataset(file: UploadFile = File(...)):
    """
    Module 1: Upload Dataset
    Reads uploaded CSV/TSV/Excel/ODS file, profiles it, and returns initial metrics.

    Loading is delegated to read_tabular(), which detects the text encoding and the
    delimiter and normalises the column labels, so non-UTF-8 exports, semicolon
    separated files and sheets with numeric headers no longer fail at the door.
    """
    contents = await file.read()

    try:
        df, load_meta = read_tabular(contents, file.filename)
    except FileLoadError as e:
        raise HTTPException(status_code=400, detail=str(e))
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Error reading file: {str(e)}")

    if len(df) == 0:
        raise HTTPException(status_code=400, detail="The file contains headers but no data rows.")

    # Store in session state
    SESSION["df"] = df
    SESSION["cleaned_df"] = None
    SESSION["suggestions"] = []
    SESSION["history"] = []
    SESSION["filename"] = file.filename

    # Module 2 & 3: Run Profiling and Statistical Anomaly Detection
    try:
        profile = profile_dataset(df)
        anomalies = detect_anomalies(df)

        SESSION["profiling"] = profile
        SESSION["anomalies"] = anomalies

        # Calculate dynamic quality score
        from backend.cleaning.validator import analyze_data_quality_metrics
        quality_metrics = analyze_data_quality_metrics(df, is_cleaned=False)
        total_cells = len(df) * len(df.columns) if len(df) > 0 else 1
        total_issues = sum(quality_metrics[k] for k in quality_metrics if k != "duplicates") + (quality_metrics["duplicates"] * len(df.columns))
        initial_score = max(10, min(100, round(100 * (1 - (total_issues / total_cells)))))

        return sanitize_for_json({
            "success": True,
            "filename": file.filename,
            "load_info": load_meta,
            "profile": profile,
            "anomalies": anomalies,
            "quality_score": initial_score,
            "preview": df_to_records(df, limit=10)
        })
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error profiling dataset: {str(e)}")

def make_ui_safe_recommendations(recommendations: dict, limit: int = 15) -> dict:
    import copy
    if not isinstance(recommendations, dict):
        return recommendations
    ui_recs = copy.deepcopy(recommendations)
    for col_rec in ui_recs.get("columns", []):
        mappings = col_rec.get("value_mappings", {})
        if isinstance(mappings, dict) and len(mappings) > limit:
            truncated = {k: mappings[k] for k in list(mappings.keys())[:limit]}
            col_rec["value_mappings"] = truncated
            col_rec["value_mappings_preview_only"] = True
            col_rec["value_mappings_total_count"] = len(mappings)
    return ui_recs

@app.post("/api/analyze")
def analyze_dataset(request: AnalyzeRequest):
    """
    Module 4 & 5: LLM Semantic Analysis & Recommendations
    Performs LLM query (or Mock fallback) and returns structured suggestions.
    """
    df = SESSION.get("df")
    if df is None:
        raise HTTPException(status_code=400, detail="No dataset uploaded. Please upload a dataset first.")
        
    profile = SESSION.get("profiling")
    anomalies = SESSION.get("anomalies")
    
    api_key = request.api_key
    if not api_key:
        if "openai" in request.provider.lower():
            api_key = os.getenv("OPENAI_API_KEY")
        elif "anthropic" in request.provider.lower():
            api_key = os.getenv("ANTHROPIC_API_KEY")
            
    try:
        recommendations, system_prompt, user_prompt = generate_recommendations(df, profile, anomalies, request.provider, api_key)
        SESSION["suggestions"] = recommendations
        SESSION["system_prompt"] = system_prompt
        SESSION["user_prompt"] = user_prompt
        
        # Prepare a UI-safe, truncated version of recommendations for the API response
        ui_recommendations = make_ui_safe_recommendations(recommendations)
        
        return sanitize_for_json({
            "success": True,
            "recommendations": ui_recommendations,
            "system_prompt": system_prompt,
            "user_prompt": user_prompt
        })
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"LLM Recommendation Error: {str(e)}")

@app.post("/api/clean")
def clean_dataset(request: CleanRequest):
    """
    Module 6: Automatic Data Cleaning Execution
    Executes Pandas operations based on recommendations and returns history log.
    """
    df = SESSION.get("df")
    if df is None:
        raise HTTPException(status_code=400, detail="No dataset uploaded.")
        
    suggestions = SESSION.get("suggestions")
    if not suggestions:
        raise HTTPException(status_code=400, detail="No recommendations found. Please run analysis first.")
        
    try:
        cleaned_df, history = apply_cleaning_recommendations(df, suggestions, request.active_cols)
        SESSION["cleaned_df"] = cleaned_df
        SESSION["history"] = history
        
        return sanitize_for_json({
            "success": True,
            "history": history,
            "preview": df_to_records(cleaned_df, limit=10)
        })
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error executing cleaning: {str(e)}")

@app.get("/api/validate")
def validate_dataset():
    """
    Module 7: Validate Clean Dataset
    Compares metrics before and after cleaning.
    """
    df = SESSION.get("df")
    cleaned_df = SESSION.get("cleaned_df")
    
    if df is None or cleaned_df is None:
        raise HTTPException(status_code=400, detail="Ensure dataset is uploaded and cleaning is applied.")
        
    try:
        validation_report = validate_cleaning(df, cleaned_df)
        return sanitize_for_json({
            "success": True,
            "validation": validation_report
        })
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Validation Error: {str(e)}")

def create_styled_xlsx(df: pd.DataFrame) -> io.BytesIO:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cleaned Dataset"
    
    # Ensure grid lines are visible
    ws.views.sheetView[0].showGridLines = True
    
    # Header styling (Sleek Slate Dark grey theme matching our academic template)
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    # Data row styling
    data_font = Font(name="Calibri", size=11, color="000000")
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    
    # Write headers
    headers = list(df.columns)
    ws.append(headers)
    ws.row_dimensions[1].height = 28
    
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align
        cell.border = thin_border
        
    # Write data rows
    phone_cols = [i for i, col in enumerate(headers) if any(x in col.lower() for x in ["phone", "mobile", "contact", "tel", "zip"])]
    date_cols = [i for i, col in enumerate(headers) if "date" in col.lower() or "admission" in col.lower() or "enrollment" in col.lower()]
    
    for row_idx, row in enumerate(df.itertuples(index=False), start=2):
        ws.row_dimensions[row_idx].height = 20
        for col_idx, val in enumerate(row, start=1):
            cell = ws.cell(row=row_idx, column=col_idx)
            
            # Format phone numbers as text to prevent scientific notation in Excel
            if (col_idx - 1) in phone_cols:
                cell.value = str(val) if pd.notna(val) else ""
                cell.number_format = '@'  # Text format
                cell.alignment = Alignment(horizontal="left", vertical="center")
            else:
                cell.value = val if pd.notna(val) else ""
                if (col_idx - 1) in date_cols:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif isinstance(val, (int, float)):
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                    
            cell.font = data_font
            cell.border = thin_border

    # Enable auto-filter
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(df) + 1}"

    # Auto-adjust column widths based on maximum string length
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer

@app.get("/api/export")
def export_dataset(format: str = "csv"):
    """
    Module 8: Export Clean Dataset
    Sends back the cleaned pandas DataFrame as a downloadable CSV, Excel, or ODS.
    """
    cleaned_df = SESSION.get("cleaned_df")
    if cleaned_df is None:
        # Fall back to original if cleaning wasn't run
        cleaned_df = SESSION.get("df")
        
    if cleaned_df is None:
        raise HTTPException(status_code=400, detail="No dataset found to export.")
        
    # Replace NaN values with empty string/None for clean Excel/ODS export if needed
    export_df = cleaned_df.copy()
    
    # 1. Standardize phone and dates inside the dataframe itself so ALL exports (CSV/ODS/Excel) match
    for col in export_df.columns:
        if any(x in col.lower() for x in ["phone", "mobile", "contact", "tel", "zip"]):
            def fmt_phone(val):
                if pd.isna(val) or str(val).strip() == "" or str(val).lower() == "nan":
                    return ""
                v_str = str(val).strip()
                if v_str.endswith(".0"):
                    v_str = v_str[:-2]
                if "e+" in v_str.lower():
                    try:
                        v_str = str(int(float(val)))
                    except:
                        pass
                return v_str
            export_df[col] = export_df[col].apply(fmt_phone)
            
        if any(x in col.lower() for x in ["date", "admission", "enrollment"]):
            def fmt_date(val):
                if pd.isna(val) or str(val).strip() == "" or str(val).lower() == "nan":
                    return ""
                val_str = str(val).strip()
                # Try common formats first (ISO formats first to prevent corruption of clean values)
                for fmt in ('%Y-%m-%d', '%Y/%m/%d', '%d-%m-%Y', '%d/%m/%Y', '%m-%d-%Y', '%m/%d/%Y'):
                    try:
                        import datetime
                        parsed_dt = datetime.datetime.strptime(val_str, fmt)
                        return parsed_dt.strftime('%Y-%m-%d')
                    except ValueError:
                        continue
                # Try general pandas dayfirst=True parsing
                try:
                    parsed = pd.to_datetime(val_str, errors='coerce', dayfirst=True)
                    if pd.notna(parsed):
                        return parsed.strftime('%Y-%m-%d')
                except:
                    pass
                return val_str
            export_df[col] = export_df[col].apply(fmt_date)
            
    try:
        if format == "xlsx":
            buffer = create_styled_xlsx(export_df)
            response = StreamingResponse(
                buffer,
                media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
            response.headers["Content-Disposition"] = "attachment; filename=cleaned_dataset.xlsx"
            return response
            
        elif format == "ods":
            buffer = io.BytesIO()
            # Save to ODS using df.to_excel with odf engine
            export_df.to_excel(buffer, index=False, engine="odf")
            buffer.seek(0)
            response = StreamingResponse(
                buffer,
                media_type="application/vnd.oasis.opendocument.spreadsheet"
            )
            response.headers["Content-Disposition"] = "attachment; filename=cleaned_dataset.ods"
            return response
            
        else: # Default is CSV
            csv_data = export_df.to_csv(index=False)
            csv_bytes = csv_data.encode('utf-8-sig')
            response = StreamingResponse(
                io.BytesIO(csv_bytes),
                media_type="text/csv"
            )
            response.headers["Content-Disposition"] = "attachment; filename=cleaned_dataset.csv"
            return response
            
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Export Error: {str(e)}")

@app.get("/api/report")
def get_cleaning_report():
    """
    Returns a professionally formatted ReportLab PDF report summarizing statistical quality issues,
    cleaning transformations applied, quality scores before vs after, and validation conclusion.
    """
    df = SESSION.get("df")
    cleaned_df = SESSION.get("cleaned_df")
    
    if df is None or cleaned_df is None:
        raise HTTPException(status_code=400, detail="Ensure dataset is uploaded and cleaning is applied.")
        
    try:
        original_profile = SESSION.get("profiling")
        original_anomalies = SESSION.get("anomalies")
        history = SESSION.get("history", [])
        validation_report = validate_cleaning(df, cleaned_df)
        filename = SESSION.get("filename", "dataset.csv")
        
        # Generate the ReportLab PDF bytes
        pdf_bytes = generate_pdf_report(filename, original_profile, original_anomalies, history, validation_report, df, cleaned_df)
        
        response = StreamingResponse(
            io.BytesIO(pdf_bytes),
            media_type="application/pdf"
        )
        response.headers["Content-Disposition"] = "attachment; filename=cleaning_report.pdf"
        return response
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error generating PDF report: {str(e)}")

# --- PHASE 2: ML PIPELINE ENDPOINTS ---

class MLUnderstandRequest(BaseModel):
    provider: str
    api_key: Optional[str] = None

class MLAnalyzeObjectiveRequest(BaseModel):
    question: str
    provider: str
    api_key: Optional[str] = None

class MLTrainRequest(BaseModel):
    target: str
    features: List[str]
    problem_type: str
    feature_selection_mode: Optional[str] = "automatic"
    excluded_features: Optional[List[str]] = None


@app.post("/api/ml/understand")
def ml_understand(request: MLUnderstandRequest):
    """
    Reads the validated cleaned dataset from Phase 1.
    If missing, returns an HTTP 400 error.
    """
    cleaned_df = SESSION.get("cleaned_df")
    if cleaned_df is None:
        raise HTTPException(status_code=400, detail="Please complete Phase 1 cleaning and validation before starting Phase 2.")
        
    try:
        from backend.cleaning.profiler import profile_dataset
        from backend.ml.question_generator import generate_questions
        
        # Profile cleaned dataset
        profile = profile_dataset(cleaned_df)
        
        # Classify columns
        numerical = []
        categorical = []
        date = []
        identifier = []
        text = []
        possible_targets = []
        
        for col in profile["columns"]:
            name = col["name"]
            ctype = col["type"]
            
            if ctype == "Numeric":
                numerical.append(name)
                possible_targets.append(name)
            elif ctype == "Categorical":
                categorical.append(name)
                possible_targets.append(name)
            elif ctype == "Boolean":
                categorical.append(name)  # Treat booleans as categoricals for ML grouping
                possible_targets.append(name)
            elif ctype == "Date":
                date.append(name)
            elif ctype == "Identifier":
                identifier.append(name)
            else: # Free Text, Email, Phone
                text.append(name)
                
        # Generate questions
        questions = generate_questions(cleaned_df, profile["columns"], request.provider, request.api_key)
        
        return sanitize_for_json({
            "success": True,
            "summary": {
                "dataset_name": SESSION.get("filename", "cleaned_dataset.csv"),
                "rows": len(cleaned_df),
                "columns": len(cleaned_df.columns),
                "numerical": numerical,
                "categorical": categorical,
                "date": date,
                "identifier": identifier,
                "text": text,
                "possible_targets": possible_targets
            },
            "questions": questions
        })
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error analyzing cleaned dataset: {str(e)}")

@app.post("/api/ml/analyze-objective")
def ml_analyze_objective(request: MLAnalyzeObjectiveRequest):
    """
    Analyzes user selected or typed question to automatically identify target,
    features, excluded columns, and task type.
    """
    cleaned_df = SESSION.get("cleaned_df")
    if cleaned_df is None:
        raise HTTPException(status_code=400, detail="No validated cleaned dataset available. Please complete Phase 1 first.")
        
    try:
        from backend.cleaning.profiler import profile_dataset
        from backend.ml.objective_analyzer import analyze_objective
        
        profile = profile_dataset(cleaned_df)
        analysis = analyze_objective(cleaned_df, profile["columns"], request.question, request.provider, request.api_key)
        
        return sanitize_for_json({
            "success": True,
            "analysis": analysis
        })
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error analyzing objective: {str(e)}")

@app.post("/api/ml/train")
def ml_train(request: MLTrainRequest):
    """
    Runs Feature Engineering, Feature Selection, Model Training, Model Evaluation,
    and Best Model Selection on the cleaned dataset.
    """
    cleaned_df = SESSION.get("cleaned_df")
    if cleaned_df is None:
        raise HTTPException(status_code=400, detail="No validated cleaned dataset available. Please complete Phase 1 first.")

    try:
        from backend.cleaning.profiler import profile_dataset
        from backend.ml.feature_engineering import FeatureEngineer
        from backend.ml.feature_selection import select_features
        from backend.ml.model_trainer import (
            train_models, cross_validate_models, train_baselines,
        )
        from backend.ml.evaluator import evaluate_all_models
        from backend.ml.model_selector import select_best_model
        from backend.ml.objective_analyzer import resolve_problem_type, _is_excluded_column, exclusion_reason
        from sklearn.model_selection import train_test_split
        import numpy as np

        # 1. Column types from the cleaned dataset
        profile = profile_dataset(cleaned_df)
        col_types = {col["name"]: col["type"] for col in profile["columns"]}

        if request.target not in cleaned_df.columns:
            raise HTTPException(status_code=400, detail=f"Target column '{request.target}' not found in dataset.")

        all_non_target = [c for c in cleaned_df.columns if c != request.target]

        # Enforce strict separation:
        # 1. Intentionally excluded by policy/identifier rules (never candidates for feature selection)
        # 2. Excluded by user selection (if unchecked in UI)
        # 3. Active eligible features to be evaluated by feature selection
        excluded_by_policy = []
        excluded_before_selection = []
        active_features = []

        for col in all_non_target:
            if _is_excluded_column(col, col_types.get(col)):
                reason = exclusion_reason(col, col_types.get(col))
                excluded_before_selection.append(col)
                excluded_by_policy.append({"feature": col, "reason": reason})
            elif col not in request.features:
                reason = "Excluded by user selection"
                excluded_before_selection.append(col)
                excluded_by_policy.append({"feature": col, "reason": reason})
            else:
                active_features.append(col)

        if not active_features:
            raise HTTPException(status_code=400, detail="At least one valid feature column must be selected.")

        # 2. The backend decides the problem type. The frontend used to derive it
        #    separately and post it back as scraped DOM text, so a continuous
        #    target could arrive labelled "Classification"; every classifier then
        #    failed with "Unknown label type: continuous" and the UI rendered a
        #    best model of "None" with blank metrics.
        problem_type, problem_reason = resolve_problem_type(
            cleaned_df[request.target], col_types.get(request.target), request.problem_type
        )

        if len(cleaned_df) < 10:
            raise HTTPException(
                status_code=400,
                detail=f"Only {len(cleaned_df)} rows remain after cleaning - too few to train and evaluate a model."
            )

        # 3. Split once, up front. Classification is stratified so every class is
        #    represented in both partitions.
        stratify = None
        if problem_type == "Classification":
            counts = cleaned_df[request.target].astype(str).value_counts()
            if len(counts) > 1 and counts.min() >= 2:
                stratify = cleaned_df[request.target].astype(str)
        train_df, test_df = train_test_split(
            cleaned_df, test_size=0.2, random_state=42, stratify=stratify
        )

        # 4. Fit preprocessing on the training partition only
        engineer = FeatureEngineer(request.target, active_features, col_types)
        engineer.fit(train_df)

        X_train, y_train, feature_names = engineer.transform(train_df)
        X_test, y_test, _ = engineer.transform(test_df)

        # 5. Feature selection on the training partition, scored one row per
        #    original column rather than per one-hot dummy
        X_select, select_names, discrete_mask = engineer.transform_for_selection(train_df)
        selection_results = select_features(
            X_select, y_train, select_names, problem_type, discrete_mask=discrete_mask
        )

        mode = getattr(request, "feature_selection_mode", "automatic") or "automatic"
        if mode == "automatic":
            keep_sources = selection_results.get("selected") or list(select_names)
        else:
            keep_sources = list(select_names)
        selection_results["applied_mode"] = mode
        selection_results["applied_features"] = keep_sources
        selection_results["excluded_before_selection"] = excluded_before_selection
        # Surfaced separately from the MI/significance ranking so the user can
        # tell "intentionally excluded by policy" (Age, Student_ID, ...) apart
        # from a column that simply never scored well enough to be selected.
        selection_results["excluded_by_policy"] = excluded_by_policy
        selection_results["target"] = request.target
        selection_results["all_columns"] = list(cleaned_df.columns)

        # Validation check: Ensure every original non-target column appears exactly once
        selected_set = set(selection_results.get("selected", []))
        removed_set = set(selection_results.get("removed", []))
        excluded_set = set(excluded_before_selection)

        assert selected_set.isdisjoint(removed_set), "Overlap detected between selected and removed features"
        assert (selected_set | removed_set) == set(active_features), "Eligible features mismatch with selected + removed"
        assert set(active_features).isdisjoint(excluded_set), "Overlap between eligible features and excluded features"
        total_accounted = selected_set | removed_set | excluded_set | {request.target}
        assert total_accounted == set(cleaned_df.columns), f"Missing columns in pipeline accounting: {set(cleaned_df.columns) - total_accounted}"

        selection_results["validation"] = {
            "is_valid": True,
            "target": request.target,
            "total_columns": len(cleaned_df.columns),
            "eligible_features_count": len(active_features),
            "selected_count": len(selection_results.get("selected", [])),
            "removed_count": len(selection_results.get("removed", [])),
            "excluded_count": len(excluded_by_policy),
        }

        keep_indices = engineer.model_columns_for(keep_sources)
        if keep_indices:
            X_train = X_train[:, keep_indices]
            X_test = X_test[:, keep_indices]
            feature_names = [feature_names[i] for i in keep_indices]

        if X_train.shape[1] == 0:
            raise HTTPException(status_code=400, detail="No usable feature columns remain after preprocessing.")

        # 6. Train candidate models on the training partition
        fitted_models, training_info = train_models(X_train, y_train, problem_type)
        if not fitted_models:
            raise HTTPException(
                status_code=500,
                detail="No candidate model could be trained on this data. Check the target and feature selection."
            )

        # 7. Evaluate on the untouched test partition
        evaluations = evaluate_all_models(
            fitted_models, X_test, y_test, problem_type,
            X_train=X_train, y_train=y_train,
            class_names=engineer.target_classes,
        )

        # 7b. Cross-validated scores, so the recommendation rests on more than one
        #     split, plus a trivial baseline to measure the models against.
        cv_results = cross_validate_models(fitted_models, X_train, y_train, problem_type)

        # Every trivial strategy is evaluated and the strongest becomes the bar the
        # models must clear. Macro-F1 rewards a balanced classifier just for
        # attempting the minority class, so comparing only against a
        # most-frequent baseline flatters a model with no signal.
        baseline_models = train_baselines(X_train, y_train, problem_type)
        baseline_metrics = {}
        if baseline_models:
            baseline_evals = evaluate_all_models(
                baseline_models, X_test, y_test, problem_type,
                X_train=X_train, y_train=y_train, class_names=engineer.target_classes,
            )
            baseline_cv = cross_validate_models(baseline_models, X_train, y_train, problem_type)

            def _bscore(name):
                if name in baseline_cv:
                    return baseline_cv[name]["mean"]
                metrics = baseline_evals.get(name, {})
                value = metrics.get("f1_macro") if problem_type == "Classification" else metrics.get("r2_score")
                return value if value is not None else float("-inf")

            strongest = max(baseline_evals, key=_bscore) if baseline_evals else None
            if strongest:
                baseline_metrics = dict(baseline_evals[strongest])
                baseline_metrics["strategy"] = strongest
                baseline_metrics["cv"] = baseline_cv.get(strongest)
                baseline_metrics["all_strategies"] = {
                    name: {"metrics": baseline_evals.get(name, {}), "cv": baseline_cv.get(name)}
                    for name in baseline_evals
                }

        # 8. Select the best model on the cross-validated evidence
        recommendation = select_best_model(
            evaluations, problem_type, request.target,
            cv_results=cv_results, baseline_metrics=baseline_metrics,
            training_info=training_info,
        )

        # 9. Feature engineering summary
        applied_transformations = []
        if engineer.onehot_sources:
            applied_transformations.append(
                f"One-hot encoding ({len(engineer.onehot_sources)} categorical column(s))"
            )
        if engineer.frequency_sources:
            applied_transformations.append(
                f"Frequency encoding ({len(engineer.frequency_sources)} high-cardinality column(s))"
            )
        if any(col_types.get(f) == "Numeric" for f in active_features):
            applied_transformations.append("Median imputation and standard scaling of numeric columns")
        if engineer.date_sources:
            applied_transformations.append(
                "Date part extraction with cyclical month/weekday encoding"
            )
        applied_transformations.append("Preprocessing fitted on the training split only")

        return sanitize_for_json({
            "success": True,
            "problem_type": problem_type,
            "problem_type_reason": problem_reason,
            "split": {
                "train_rows": int(len(train_df)),
                "test_rows": int(len(test_df)),
                "stratified": stratify is not None,
            },
            "feature_engineering": {
                "transformations": applied_transformations,
                "engineered_cols": feature_names,
            },
            "feature_selection": selection_results,
            "models_trained": list(fitted_models.keys()),
            "evaluations": evaluations,
            "cross_validation": cv_results,
            "baseline": baseline_metrics,
            "training_info": training_info,
            "best_model": recommendation,
        })
    except HTTPException:
        raise
    except Exception as e:
        import traceback
        tb = traceback.format_exc()
        print("--- PIPELINE TRAINING EXCEPTION ---")
        print(tb)
        raise HTTPException(status_code=500, detail=f"Error running pipeline training: {str(e)}")

# Mount frontend files at project root
# Ensure frontend directory exists before running
os.makedirs("frontend", exist_ok=True)
app.mount("/", StaticFiles(directory="frontend", html=True), name="static")
